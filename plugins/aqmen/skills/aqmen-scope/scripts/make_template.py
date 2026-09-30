#!/usr/bin/env python
"""
(Re)build assets/scope-template.docx: the empty, branded Word shell that
build_scope.py fills. Run only when the brand or page set-up changes.

    python make_template.py <source.docx>

The source is any past Aqmen proposal in Word (it supplies page size, margins,
header/footer and list numbering). This script strips the body and re-brands the
styles: Montserrat, aqmen navy/ink/grey, brand border colours. The template is
intentionally blank when opened — its job is to carry styles, not content.
"""
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "assets" / "scope-template.docx"

FONT = "Montserrat"
INK = "302E2E"
NAVY = "03045E"
GREY = "8A8A8A"
BORDER = "D6EEF5"


def set_font(rPr_owner_font, name=FONT):
    rPr_owner_font.name = name
    rpr = rPr_owner_font.element.rPr if hasattr(rPr_owner_font, "element") else None


def brand_style(style, size=None, color=INK, bold=None):
    f = style.font
    f.name = FONT
    # python-docx sets ascii/hAnsi; also set eastAsia/cs so Word never falls back
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rFonts.get(qn(attr)) is not None:
            del rFonts.attrib[qn(attr)]
    if size:
        f.size = Pt(size)
    if color:
        f.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        f.bold = bold


def rebrand_theme(doc):
    """Point the theme's major/minor fonts at Montserrat so anything that still
    inherits from the theme (tables, fields) follows the brand."""
    for part in doc.part.package.parts:
        if part.partname.endswith("theme/theme1.xml"):
            xml = part.blob.decode("utf-8")
            import re
            xml = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*(")', r"\g<1>" + FONT + r"\2", xml)
            part._blob = xml.encode("utf-8")


def recolor_runs(paragraph, color_map):
    for r in paragraph.runs:
        if r.font.color is not None and r.font.color.rgb is not None:
            hexv = str(r.font.color.rgb)
            if hexv in color_map:
                r.font.color.rgb = RGBColor.from_string(color_map[hexv])
        r.font.name = FONT
        rPr = r._r.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts"); rPr.insert(0, rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rFonts.set(qn(attr), FONT)


def main(src):
    doc = Document(src)
    body = doc.element.body
    for el in list(body):
        if not el.tag.endswith("}sectPr"):
            body.remove(el)

    # Styles
    brand_style(doc.styles["Normal"], size=10, color=INK)
    for name in ("List Bullet", "List Number", "Header", "Footer", "Table Grid"):
        try:
            brand_style(doc.styles[name], color=None)
        except KeyError:
            pass
    rebrand_theme(doc)

    # Header / footer: keep layout, re-colour to brand, re-font
    legacy = {"0728A4": NAVY, "999999": GREY, "555555": GREY}
    sec = doc.sections[0]
    for p in list(sec.header.paragraphs) + list(sec.footer.paragraphs):
        recolor_runs(p, legacy)
        pPr = p._p.pPr
        if pPr is not None:
            for b in pPr.iter(qn("w:bottom")):
                b.set(qn("w:color"), BORDER)

    doc.core_properties.title = ""
    doc.core_properties.subject = ""
    doc.core_properties.author = "Aqmen AI"
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    main(sys.argv[1])
