#!/usr/bin/env python
"""
Render a storyline plan (one row per slide) from JSON into:

    <slug> - Storyline plan.xlsx   the plan: Plan sheet (one row per slide),
                                   Storyline sheet (the action titles in
                                   reading order), Data request sheet

Usage:
    python build_plan.py plan.json <outdir>

The ghost deck is not built here: it is built in the aqmen workspace with
create_deck (see ../SKILL.md). The plan schema is documented in
../references/plan-format.md and exemplified in
../assets/example-carwash-plan.json. Requires openpyxl.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path


COLUMNS = [
    ("#", 5), ("Section", 16), ("Slide type", 12), ("Title (action title)", 46),
    ("Content", 46), ("Exhibit / analysis", 30), ("Data & sources", 28),
    ("Owner", 11), ("Purpose (so-what)", 36), ("Image prompt", 44), ("Status", 10),
]
SLIDE_TYPES = ("cover", "agenda", "divider", "exec_summary", "content", "appendix")


def is_appendix(name):
    return str(name).strip().lower() == "appendix"


def load_plan(path):
    plan = json.loads(Path(path).read_text(encoding="utf-8"))
    plan.setdefault("slug", re.sub(r"[^A-Za-z0-9]+", "-", plan["title"]).strip("-"))
    for i, row in enumerate(plan["slides"], 1):
        row.setdefault("n", i)
        row.setdefault("type", "content")
        row.setdefault("status", "planned")
        for k in ("section", "title", "content", "exhibit", "data", "owner", "purpose", "image_prompt"):
            row.setdefault(k, "")
    return plan


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def build_xlsx(plan, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    NAVY, BLUE, INK, GREY, FILL, WHITE = "03045E", "0728A3", "302E2E", "8A8A8A", "F4F9FC", "FFFFFF"
    thin = Side(style="thin", color=NAVY)
    hair = Side(style="hair", color="C9CCD6")

    wb = Workbook()
    ws = wb.active
    ws.title = "Plan"

    # title block
    ws["A1"] = plan["title"]
    ws["A1"].font = Font(name="Montserrat", size=16, bold=True, color=NAVY)
    ws["A2"] = plan.get("subtitle", "Storyline plan — one row per slide")
    ws["A2"].font = Font(name="Montserrat", size=10, color=GREY)
    meta = " · ".join(x for x in (plan.get("client"), plan.get("date"), f"{len(plan['slides'])} slides") if x)
    ws["A3"] = meta
    ws["A3"].font = Font(name="Montserrat", size=9, color=GREY)

    head_row = 5
    for j, (name, width) in enumerate(COLUMNS, 1):
        c = ws.cell(row=head_row, column=j, value=name.upper())
        c.font = Font(name="Montserrat", size=9, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[head_row].height = 28

    r = head_row + 1
    current_section = None
    for row in plan["slides"]:
        is_break = row["type"] in ("cover", "agenda", "divider", "exec_summary", "appendix")
        vals = [row["n"], row["section"], row["type"], row["title"], row["content"], row["exhibit"],
                row["data"], row["owner"], row["purpose"], row["image_prompt"], row["status"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = Font(name="Montserrat", size=9, color=INK,
                          bold=(j == 4 and not is_break) or (is_break and j in (2, 3, 4)))
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=hair, left=hair, right=hair)
            if is_break:
                c.fill = PatternFill("solid", fgColor="D6EEF5")
            elif row["section"] != current_section and j == 2:
                c.font = Font(name="Montserrat", size=9, color=BLUE, bold=True)
        # row height from the longest cell (~ chars per line at the column width)
        # Montserrat 9pt fits roughly one character per unit of column width
        longest = max((len(str(v)) / max(COLUMNS[j][1] - 2, 4) for j, v in enumerate(vals)), default=1)
        ws.row_dimensions[r].height = max(18, min(170, 12.5 * (int(longest) + 1)))
        current_section = row["section"]
        r += 1

    ws.freeze_panes = ws.cell(row=head_row + 1, column=5)
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(COLUMNS))}{r - 1}"
    ws.sheet_view.zoomScale = 90

    # second sheet: storyline in reading order (section → titles), the "dash" view
    ws2 = wb.create_sheet("Storyline")
    ws2["A1"] = "Storyline — read the action titles top to bottom; they must argue the case on their own"
    ws2["A1"].font = Font(name="Montserrat", size=11, bold=True, color=NAVY)
    ws2.column_dimensions["A"].width = 6
    ws2.column_dimensions["B"].width = 18
    ws2.column_dimensions["C"].width = 100
    rr = 3
    for row in plan["slides"]:
        if row["type"] in ("cover", "agenda"):
            continue
        ws2.cell(row=rr, column=1, value=row["n"]).font = Font(name="Montserrat", size=9, color=GREY)
        ws2.cell(row=rr, column=2, value=row["section"]).font = Font(name="Montserrat", size=9, color=BLUE, bold=True)
        c = ws2.cell(row=rr, column=3, value=row["title"])
        c.font = Font(name="Montserrat", size=10, color=INK,
                      bold=row["type"] in ("divider", "appendix", "exec_summary"))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        rr += 1

    # third sheet: data request derived from the plan
    ws3 = wb.create_sheet("Data request")
    heads = ("#", "Slide", "Data & sources", "Owner", "Status")
    for j, h in enumerate(heads, 1):
        c = ws3.cell(row=1, column=j, value=h.upper())
        c.font = Font(name="Montserrat", size=9, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
    for j, w in enumerate((6, 50, 60, 14, 12), 1):
        ws3.column_dimensions[get_column_letter(j)].width = w
    rr = 2
    for row in plan["slides"]:
        if not row["data"]:
            continue
        for j, v in enumerate((row["n"], row["title"], row["data"], row["owner"], "open"), 1):
            c = ws3.cell(row=rr, column=j, value=v)
            c.font = Font(name="Montserrat", size=9, color=INK)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        rr += 1

    wb.save(path)
    return path


def main(argv):
    if len(argv) < 3:
        raise SystemExit(__doc__)
    plan = load_plan(argv[1])
    outdir = Path(argv[2]); outdir.mkdir(parents=True, exist_ok=True)
    p = outdir / f"{plan['slug']} - Storyline plan.xlsx"
    build_xlsx(plan, str(p)); print(f"xlsx  {p}")


if __name__ == "__main__":
    main(sys.argv)
