#!/usr/bin/env python
"""
Render a business-plan assessment from JSON into:

    <slug> - BP Assessment.pptx   deck slides in the house style: context (growth
                                  channels, bridges), sensitivity, the assessment
                                  matrix, one deep dive per assessed assumption,
                                  Aqmen's view
    <slug> - BP Assessment.xlsx   the assessment table: every assumption with its
                                  plan values, the three lenses, the rating, the
                                  rationale, Aqmen's view and the sensitivity

Usage:
    python build_bp.py assessment.json <outdir> [--only deck|xlsx] [--final]

Schema: ../references/bp-format.md. Example: ../assets/example-wealth-bp.json.
Requires python-pptx and openpyxl.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
REFS = SKILL / "references"
sys.path.insert(0, str(REFS))

# Rating scale, in the reference deck's legend order and colours.
RATINGS = {
    "highly_conservative": ("++", "03045E", "FFFFFF", "Highly Conservative"),
    "conservative":        ("+",  "BAC8F8", "03045E", "Conservative"),
    "realistic":           ("=",  "00A650", "FFFFFF", "Realistic"),
    "optimistic":          ("-",  "FFC000", "302E2E", "Optimistic"),
    "highly_optimistic":   ("--", "C00000", "FFFFFF", "Highly Optimistic"),
    "na":                  ("N/A", "969696", "FFFFFF", "Not Applicable"),
}
LENSES = (("market", "Market"), ("competitive", "Competitive position"), ("track_record", "Track record"))


def norm_rating(v):
    if not v:
        return "na"
    k = str(v).strip().lower().replace(" ", "_").replace("-", "_")
    aliases = {"hc": "highly_conservative", "c": "conservative", "r": "realistic", "o": "optimistic",
               "ho": "highly_optimistic", "n/a": "na", "not_applicable": "na", "++": "highly_conservative",
               "+": "conservative", "=": "realistic", "-": "optimistic", "--": "highly_optimistic"}
    return aliases.get(k, k if k in RATINGS else "na")


def load(path):
    a = json.loads(Path(path).read_text(encoding="utf-8"))
    a.setdefault("slug", re.sub(r"[^A-Za-z0-9]+", "-", a["title"]).strip("-"))
    a.setdefault("years", [])
    a.setdefault("metric", "EBITDA")
    n = 0
    for cat in a["categories"]:
        for asm in cat["assumptions"]:
            n += 1
            asm.setdefault("id", str(n))
            asm["rating"] = norm_rating(asm.get("rating"))
            asm.setdefault("values", {})
            asm.setdefault("lenses", {})
            for k, _ in LENSES:
                l = asm["lenses"].setdefault(k, {})
                l["rating"] = norm_rating(l.get("rating"))
                l.setdefault("note", "")
            asm.setdefault("rationale", "")
            asm.setdefault("aqmen_view", "")
    return a


def all_assumptions(a):
    for cat in a["categories"]:
        for asm in cat["assumptions"]:
            yield cat, asm


# --------------------------------------------------------------------------- #
# Deck
# --------------------------------------------------------------------------- #
def build_deck(a, path, draft=True):
    from aqmen_deck import (Deck, Bullet, ChartSpec, Palette, Font, MARGIN, HEADER_ROW, RULE_Y,
                            CONTENT_TOP, CONTENT_BOTTOM, RAIL_LINE_X, SLIDE_W, _c)
    from pptx.util import Inches, Pt, Emu
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    d = Deck(draft=draft)
    metric = a["metric"]
    years = a["years"]
    client = a.get("client", "[Client]")

    def tag(slide, x, y, rating, size=Inches(0.18)):
        """The rating square with its symbol (legend colours)."""
        sym, fill, ink, _ = RATINGS[norm_rating(rating)]
        sq = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, size, size)
        sq.fill.solid(); sq.fill.fore_color.rgb = _c(fill)
        sq.line.fill.background(); sq.shadow.inherit = False; d._no_style(sq)
        d._text(slide, x, y, size, size, sym, size=7 if len(sym) > 2 else 8, color=_c(ink), font=Font.HEAD,
                bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, autofit=False)

    def legend(slide):
        x = Inches(1.75)  # right of the DRAFT tag
        for key in ("highly_conservative", "conservative", "realistic", "optimistic", "highly_optimistic", "na"):
            tag(slide, x, Inches(0.13), key, Inches(0.17))
            lbl = RATINGS[key][3]
            d._text(slide, Emu(x + Inches(0.22)), Inches(0.12), Inches(1.3), Inches(0.2), lbl, size=8,
                    color=Palette.INK, font=Font.BODY, anchor=MSO_ANCHOR.MIDDLE, autofit=False)
            x = Emu(x + Inches(0.22) + Inches(0.075 * len(lbl) + 0.35))

    # --- cover + agenda -----------------------------------------------------
    d.title_slide(a["title"], a.get("subtitle", "Business plan assessment"), a.get("date"))
    parts = ["Context", "Sensitivity", "Assessment", "Deep dives", "Aqmen view"]
    d.agenda(parts)

    # --- context: growth channels --------------------------------------------
    if a.get("channels"):
        d.agenda(parts, active="Context")
        ch = a["channels"]
        metrics = list(ch.get("metrics", []))
        columns = ["Metric"] + [c["name"] for c in ch["items"]]
        rows = [["Group"] + [c.get("group", "") for c in ch["items"]],
                ["Overview"] + [c.get("overview", "") for c in ch["items"]]]
        for m in metrics:
            rows.append([m] + [c.get("values", {}).get(m, "n/a") for c in ch["items"]])
        d.table_slide(ch.get("headline", f"{client}'s growth channels grouped for assessment"),
                      columns, rows, eyebrow=("Business Plan", "Context"), source=ch.get("source", f"{client} business plan"),
                      chart_title=ch.get("chart_title", "Growth channels and their contribution to the plan"),
                      takeaways=ch.get("so_whats"), col_align={j: "c" for j in range(1, len(columns))})

    # --- context: bridges ---------------------------------------------------
    for b in a.get("bridges", []):
        d.waterfall_slide(b["headline"], b["steps"], takeaways=b.get("so_whats"),
                          eyebrow=("Business Plan", "Context"), source=b.get("source", f"{client} business plan"),
                          chart_title=b.get("chart_title", f"{b.get('metric', metric)} bridge, {b.get('unit', '')}").strip(", "),
                          number_format=b.get("number_format", "#,##0.0"))

    # --- sensitivity ---------------------------------------------------------
    sens = [(cat, asm) for cat, asm in all_assumptions(a) if asm.get("sensitivity")]
    if sens:
        d.agenda(parts, active="Sensitivity")
        sens.sort(key=lambda t: -abs(float(t[1]["sensitivity"].get("delta", 0) or 0)))
        y0 = years[-1] if years else ""
        cols = ["#", "Assumption", f"Δ {y0} {metric}", f"% of {y0} {metric}", "From", "To", "Comment"]
        rows = []
        for cat, asm in sens:
            s_ = asm["sensitivity"]
            delta = s_.get("delta")
            rows.append([asm["id"], f"{asm['name']} ({cat['name']})",
                         (f"{delta:,.0f}" if isinstance(delta, (int, float)) else str(delta)),
                         (f"{s_['pct']:.1f}%" if isinstance(s_.get("pct"), (int, float)) else str(s_.get("pct", ""))),
                         s_.get("from", ""), s_.get("to", ""), s_.get("comment", "")])
        d.table_slide(a.get("sensitivity_headline", f"A ±10% change in each assumption shows where {metric} is most exposed"),
                      cols, rows, eyebrow=("Business Plan", "Sensitivity"),
                      source=f"Aqmen analysis based on {client}'s business plan",
                      chart_title=f"Sensitivity: change in {y0} {metric} with a 10% downside on each assumption",
                      widths=[0.5, 2.2, 1.2, 1.2, 1.0, 1.0, 3.6], col_align={2: "r", 3: "r", 4: "c", 5: "c"})

    # --- assessment matrix --------------------------------------------------
    d.agenda(parts, active="Assessment")
    slide = d._slide()
    d._page += 1
    d._chrome(slide, eyebrow=("Business Plan", "Assessment"), source=f"Aqmen analysis based on {client}'s business plan")
    legend(slide)
    overall = a.get("overall", {})
    d._headline(slide, overall.get("headline", f"Assessment of key assumptions implies the business plan is {RATINGS[norm_rating(overall.get('rating'))][3].lower()}"))
    if any(asm.get("deep_dive") for _, asm in all_assumptions(a)):
        sq = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.66), Inches(1.53), Inches(0.38), Inches(0.17))
        sq.fill.solid(); sq.fill.fore_color.rgb = Palette.LAVENDER; sq.line.fill.background(); d._no_style(sq)
        d._text(slide, Inches(11.08), Inches(1.47), Inches(1.6), Inches(0.29), "Deep dive to follow", size=11,
                color=Palette.NAVY_DEEP, font=Font.BODY, italic=True, autofit=False)

    # geometry from the reference slide
    x_cat, w_cat = Inches(0.74), Inches(1.33)
    x_vline = Inches(2.04)
    x_name, w_name = Inches(2.19), Inches(1.39)
    x_sep = Inches(3.67)
    x_vals = Inches(3.87)
    w_val, gap_val = Inches(0.58), Inches(0.23)
    n_vals = len(years)
    x_tag = Emu(x_vals + (w_val + gap_val) * n_vals + Inches(0.35))
    x_rat = Emu(x_tag + Inches(0.6))
    w_rat = Emu(SLIDE_W - Inches(0.68) - x_rat)
    top = Inches(2.14)
    total_rows = sum(len(c["assumptions"]) for c in a["categories"])
    row_h = Emu(int(min(int(Inches(0.52)), int((Inches(6.78) - top) / max(total_rows, 1)))))
    # column headers
    for i, yv in enumerate(years):
        d._text(slide, Emu(x_vals + (w_val + gap_val) * i), Inches(1.80), w_val, Inches(0.25), str(yv), size=12,
                color=Palette.NAVY_DEEP, font=Font.HEAD, bold=True, align=PP_ALIGN.CENTER, autofit=False)
    d._text(slide, Emu(x_tag - Inches(0.45)), Inches(1.73), Inches(1.3), Inches(0.38), [[("BP", {})], [("Assessment", {})]],
            size=12, color=Palette.NAVY_DEEP, font=Font.HEAD, bold=True, align=PP_ALIGN.CENTER, line_spacing=1.0,
            space_after=0, autofit=False)
    d._text(slide, x_rat, Inches(1.80), Inches(2), Inches(0.25), "Rationale", size=12, color=Palette.NAVY_DEEP,
            font=Font.HEAD, bold=True, autofit=False)
    d._line(slide, x_cat, top, Inches(11.91), color=Palette.NAVY_DEEP, weight=0.75)
    y = top
    n_cat = 0
    for cat in a["categories"]:
        n_cat += 1
        m = len(cat["assumptions"])
        cat_h = Emu(row_h * m)
        d._text(slide, x_cat, Emu(y + cat_h // 2 - Inches(0.19)), w_cat, Inches(0.38), cat["name"], size=12,
                color=Palette.NAVY_DEEP, font=Font.HEAD, bold=True, anchor=MSO_ANCHOR.MIDDLE, autofit=False)
        d._conn(slide, x_vline, Emu(y + Inches(0.1)), x_vline, Emu(y + cat_h - Inches(0.1)), color=Palette.NAVY_DEEP, weight=0.75)
        d._num_badge(slide, Emu(x_vline - Inches(0.08)), Emu(y + cat_h // 2 - Inches(0.08)), Inches(0.16), str(n_cat))
        for k, asm in enumerate(cat["assumptions"]):
            if k % 2 == 0:
                band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x_vline, y, Emu(SLIDE_W - Inches(0.68) - x_vline), row_h)
                band.fill.solid(); band.fill.fore_color.rgb = _c("EEEFFB"); band.line.fill.background(); d._no_style(band)
            d._text(slide, x_name, Emu(y + Inches(0.06)), w_name, Emu(row_h - Inches(0.1)), asm["name"], size=11,
                    color=Palette.NAVY_DEEP, font=Font.HEAD, bold=True, anchor=MSO_ANCHOR.MIDDLE, autofit=False)
            d._conn(slide, x_sep, Emu(y + Inches(0.1)), x_sep, Emu(y + row_h - Inches(0.1)), color=Palette.GREY_LINE, weight=0.5)
            for i, yv in enumerate(years):
                bx = Emu(x_vals + (w_val + gap_val) * i)
                by = Emu(y + row_h // 2 - Inches(0.165))
                box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, bx, by, w_val, Inches(0.33))
                box.fill.solid(); box.fill.fore_color.rgb = _c("F2F2F2")
                box.line.color.rgb = Palette.GREY_LINE; box.line.width = Pt(0.5); box.shadow.inherit = False; d._no_style(box)
                d._text(slide, bx, by, w_val, Inches(0.33), str(asm["values"].get(str(yv), asm["values"].get(yv, "n/a"))),
                        size=10, color=Palette.NAVY_DEEP, font=Font.BODY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, autofit=False)
            tag(slide, x_tag, Emu(y + row_h // 2 - Inches(0.09)), asm["rating"])
            if asm.get("deep_dive"):
                dd = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x_tag - Inches(0.05)), Emu(y + row_h // 2 - Inches(0.14)),
                                            Inches(0.28), Inches(0.05))
                dd.fill.solid(); dd.fill.fore_color.rgb = Palette.LAVENDER; dd.line.fill.background(); d._no_style(dd)
            rat = asm["rationale"] + ("  Not assessed in detail." if asm["rating"] == "na" else "")
            d._text(slide, x_rat, Emu(y + Inches(0.04)), w_rat, Emu(row_h - Inches(0.06)), rat,
                    size=8.5 if row_h >= Inches(0.5) else 7.5, color=Palette.INK, font=Font.BODY,
                    anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0, space_after=0, autofit=False)
            y = Emu(y + row_h)
        d._line(slide, x_cat, y, Inches(11.91), color=Palette.GREY_LINE if n_cat < len(a["categories"]) else Palette.NAVY_DEEP,
                weight=0.5 if n_cat < len(a["categories"]) else 0.75)

    # --- deep dives ----------------------------------------------------------
    dives = [(cat, asm) for cat, asm in all_assumptions(a) if asm.get("deep_dive")]
    if dives:
        d.agenda(parts, active="Deep dives")
    for cat, asm in dives:
        dd = asm["deep_dive"]
        slide = d._slide()
        d._page += 1
        d._chrome(slide, eyebrow=("BP Assessment", asm["name"]), source=dd.get("source", "Aqmen analysis"))
        legend(slide)
        d._headline(slide, dd.get("headline", f"{asm['id']}. {asm['name']}"))
        panels = dd.get("panels", [])[:2]
        panel_w = Inches(3.42)
        xs = [Inches(0.69), Inches(4.72)] if len(panels) == 2 else [Inches(0.69)]
        if len(panels) == 1:
            panel_w = Inches(7.45)
        d._line(slide, Inches(0.65), RULE_Y, Inches(12.04), color=Palette.NAVY_DEEP, weight=0.75)
        for i, pn in enumerate(panels):
            px = xs[i]
            d._text(slide, px, HEADER_ROW, Emu(panel_w - Inches(0.3)), Inches(0.24), pn.get("title", ""), size=14,
                    color=Palette.NAVY_DEEP, font=Font.HEAD, bold=True)
            if pn.get("rating"):
                tag(slide, Emu(px + panel_w - Inches(0.25)), Emu(HEADER_ROW + Inches(0.03)), pn["rating"], Inches(0.17))
            ty = CONTENT_TOP
            if pn.get("body"):
                body = pn["body"] if isinstance(pn["body"], list) else [pn["body"]]
                d._text(slide, px, ty, panel_w, Inches(1.6), [[(t, {})] for t in body], size=9, color=Palette.INK,
                        font=Font.BODY, line_spacing=1.05, space_after=4)
                ty = Emu(ty + Inches(0.22) * sum(max(1, len(t) // 60 + 1) for t in body) + Inches(0.2))
            if pn.get("chart"):
                c = pn["chart"]
                if pn.get("chart_title"):
                    d._text(slide, px, ty, panel_w, Inches(0.22), pn["chart_title"], size=10, color=Palette.NAVY_DEEP,
                            font=Font.HEAD, italic=True)
                    ty = Emu(ty + Inches(0.3))
                spec = ChartSpec(kind=c.get("kind", "column"), categories=list(c["categories"]),
                                 series=[(n, list(v)) for n, v in c["series"]],
                                 number_format=c.get("number_format", "#,##0.0"), legend=c.get("legend", True))
                cy = max(ty, Inches(3.6))
                d._add_chart(slide, spec, px, cy, panel_w, Emu(CONTENT_BOTTOM - cy))
            if i == 0 and len(panels) == 2:
                d._operator_glyph(slide, Inches(4.42), Emu(HEADER_ROW + Inches(0.12)), "+")
        assess = dd.get("assessment", {})
        # rail with the assessment title and its rating tag
        d._rail(slide, assess.get("bullets", [asm["rationale"]]), title=assess.get("title", f"{asm['name']} assessment"))
        tag(slide, Inches(12.45), Emu(HEADER_ROW + Inches(0.03)), assess.get("rating", asm["rating"]), Inches(0.17))

    # --- Aqmen view ------------------------------------------------------------
    views = [(cat, asm) for cat, asm in all_assumptions(a) if asm.get("aqmen_view")]
    if views:
        d.agenda(parts, active="Aqmen view")
        y_last = years[-1] if years else ""
        cols = ["#", "Assumption", f"Plan {y_last}", "Aqmen view", "Rating", "Why"]
        rows = [[asm["id"], f"{asm['name']} ({cat['name']})", str(asm["values"].get(str(y_last), "")),
                 asm["aqmen_view"], RATINGS[asm["rating"]][3], asm["rationale"]] for cat, asm in views]
        d.table_slide(a.get("view_headline", f"Aqmen's view of the assumptions that drive {metric}"), cols, rows,
                      eyebrow=("Business Plan", "Aqmen view"), source="Aqmen analysis",
                      chart_title="Plan assumption vs Aqmen's view, with the assessment",
                      widths=[0.4, 2.0, 1.0, 1.3, 1.3, 4.0], col_align={2: "c", 3: "c", 4: "c"})

    d.save(path)
    return len(d.prs.slides._sldIdLst)


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def build_xlsx(a, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    NAVY, INK, GREY, WHITE = "03045E", "302E2E", "8A8A8A", "FFFFFF"
    thin = Side(style="thin", color=NAVY); hair = Side(style="hair", color="C9CCD6")
    wb = Workbook(); ws = wb.active; ws.title = "Assessment"
    ws["A1"] = a["title"]; ws["A1"].font = Font(name="Montserrat", size=16, bold=True, color=NAVY)
    ws["A2"] = a.get("subtitle", "Business plan assessment — one row per assumption")
    ws["A2"].font = Font(name="Montserrat", size=10, color=GREY)
    years = a["years"]
    cols = [("#", 5), ("Category", 18), ("Assumption", 22)] + [(f"Plan {y}", 11) for y in years] + \
           [("BP assessment", 16), ("Market lens", 14), ("Market note", 34), ("Competitive lens", 14), ("Competitive note", 34),
            ("Track-record lens", 14), ("Track-record note", 34), ("Rationale", 50), ("Aqmen view", 24),
            (f"Sensitivity Δ {years[-1] if years else ''} {a['metric']}", 16), ("Sensitivity comment", 36), ("Deep dive", 10)]
    hr = 4
    for j, (name, w) in enumerate(cols, 1):
        c = ws.cell(row=hr, column=j, value=name.upper())
        c.font = Font(name="Montserrat", size=9, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY); c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[hr].height = 30
    r = hr + 1

    def rating_cell(c, key):
        sym, fill, ink, label = RATINGS[norm_rating(key)]
        c.value = f"{label} ({sym})"
        c.fill = PatternFill("solid", fgColor=fill)
        c.font = Font(name="Montserrat", size=9, bold=True, color=ink)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for cat, asm in all_assumptions(a):
        s_ = asm.get("sensitivity", {}) or {}
        vals = [asm["id"], cat["name"], asm["name"]] + [asm["values"].get(str(y), asm["values"].get(y, "")) for y in years]
        vals += [None, None, asm["lenses"]["market"]["note"], None, asm["lenses"]["competitive"]["note"],
                 None, asm["lenses"]["track_record"]["note"], asm["rationale"], asm["aqmen_view"],
                 s_.get("delta", ""), s_.get("comment", ""), "yes" if asm.get("deep_dive") else ""]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = Font(name="Montserrat", size=9, color=INK, bold=(j == 3))
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=hair, left=hair, right=hair)
        base = 3 + len(years)
        rating_cell(ws.cell(row=r, column=base + 1), asm["rating"])
        rating_cell(ws.cell(row=r, column=base + 2), asm["lenses"]["market"]["rating"])
        rating_cell(ws.cell(row=r, column=base + 4), asm["lenses"]["competitive"]["rating"])
        rating_cell(ws.cell(row=r, column=base + 6), asm["lenses"]["track_record"]["rating"])
        ws.row_dimensions[r].height = max(30, min(150, 12.5 * (len(asm["rationale"]) // 48 + 2)))
        r += 1
    ws.freeze_panes = ws.cell(row=hr + 1, column=4)
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(cols))}{r - 1}"

    ws2 = wb.create_sheet("Scale")
    ws2["A1"] = "Rating scale"; ws2["A1"].font = Font(name="Montserrat", size=11, bold=True, color=NAVY)
    ws2.column_dimensions["A"].width = 24; ws2.column_dimensions["B"].width = 8; ws2.column_dimensions["C"].width = 90
    meaning = {
        "highly_conservative": "The plan assumes materially less than the evidence supports on all three lenses.",
        "conservative": "The plan sits below what market, position and track record would justify.",
        "realistic": "In line with the market, the company's position and its own history.",
        "optimistic": "Above what at least one lens supports; achievable with execution the company has not yet shown.",
        "highly_optimistic": "Beyond what the market, the position or the track record supports; needs a step-change.",
        "na": "Not assessed in detail; carried at plan value.",
    }
    for i, (k, (sym, fill, ink, label)) in enumerate(RATINGS.items(), 3):
        c = ws2.cell(row=i, column=1, value=label); rating_cell(c, k)
        ws2.cell(row=i, column=2, value=sym).font = Font(name="Montserrat", size=9, color=INK)
        ws2.cell(row=i, column=3, value=meaning[k]).font = Font(name="Montserrat", size=9, color=INK)
    wb.save(path)
    return path


def main(argv):
    if len(argv) < 3:
        raise SystemExit(__doc__)
    a = load(argv[1])
    outdir = Path(argv[2]); outdir.mkdir(parents=True, exist_ok=True)
    only = argv[argv.index("--only") + 1] if "--only" in argv else None
    final = "--final" in argv
    base = a["slug"]
    if only in (None, "deck"):
        p = outdir / f"{base} - BP Assessment.pptx"
        n = build_deck(a, str(p), draft=not final); print(f"deck  {p}  ({n} slides)")
    if only in (None, "xlsx"):
        p = outdir / f"{base} - BP Assessment.xlsx"
        build_xlsx(a, str(p)); print(f"xlsx  {p}")


if __name__ == "__main__":
    main(sys.argv)
