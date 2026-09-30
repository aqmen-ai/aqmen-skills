#!/usr/bin/env python
"""
Render a Dot-Dash (presentation plan) from JSON into:

    <slug> - Dot-Dash.xlsx        the plan: one row per slide, house-styled Excel
    <slug> - Skeleton deck.pptx   the same plan as a skeleton deck (title, exhibit
                                  title, a placeholder describing the analysis,
                                  image prompt and purpose in the speaker notes)

Usage:
    python build_dotdash.py plan.json <outdir> [--only xlsx|deck] [--final]
    python build_dotdash.py --from-cdd content.json plan.json   # derive a plan from a cdd-output content file

The plan schema is documented in ../references/dotdash-format.md and
exemplified in ../assets/example-carwash-dotdash.json.
Requires openpyxl and python-pptx.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
REFS = SKILL / "references"
sys.path.insert(0, str(REFS))

COLUMNS = [
    ("#", 5), ("Section", 16), ("Slide type", 12), ("Title (action title)", 46),
    ("Content", 46), ("Exhibit / analysis", 30), ("Data & sources", 28),
    ("Owner", 11), ("Purpose (so-what)", 36), ("Image prompt", 44), ("Status", 10),
]
SLIDE_TYPES = ("cover", "agenda", "divider", "exec_summary", "content", "appendix")


def load_plan(path):
    plan = json.loads(Path(path).read_text(encoding="utf-8"))
    plan.setdefault("slug", re.sub(r"[^A-Za-z0-9]+", "-", plan["title"]).strip("-"))
    for i, row in enumerate(plan["slides"], 1):
        row.setdefault("n", i)
        row.setdefault("type", "content")
        row.setdefault("status", "planned")
        for k in ("section", "title", "content", "exhibit", "data", "owner", "purpose", "image_prompt"):
            row.setdefault(k, "")
    return plan


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def build_xlsx(plan, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    NAVY, BLUE, INK, GREY, FILL, WHITE = "03045E", "0728A3", "302E2E", "8A8A8A", "F4F9FC", "FFFFFF"
    thin = Side(style="thin", color=NAVY)
    hair = Side(style="hair", color="C9CCD6")

    wb = Workbook()
    ws = wb.active
    ws.title = "Dot-Dash"

    # title block
    ws["A1"] = plan["title"]
    ws["A1"].font = Font(name="Montserrat", size=16, bold=True, color=NAVY)
    ws["A2"] = plan.get("subtitle", "Presentation plan (Dot-Dash) — one row per slide")
    ws["A2"].font = Font(name="Montserrat", size=10, color=GREY)
    meta = " · ".join(x for x in (plan.get("client"), plan.get("date"), f"{len(plan['slides'])} slides") if x)
    ws["A3"] = meta
    ws["A3"].font = Font(name="Montserrat", size=9, color=GREY)

    head_row = 5
    for j, (name, width) in enumerate(COLUMNS, 1):
        c = ws.cell(row=head_row, column=j, value=name.upper())
        c.font = Font(name="Montserrat", size=9, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[head_row].height = 28

    r = head_row + 1
    current_section = None
    for row in plan["slides"]:
        is_break = row["type"] in ("cover", "agenda", "divider", "exec_summary", "appendix")
        vals = [row["n"], row["section"], row["type"], row["title"], row["content"], row["exhibit"],
                row["data"], row["owner"], row["purpose"], row["image_prompt"], row["status"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = Font(name="Montserrat", size=9, color=INK,
                          bold=(j == 4 and not is_break) or (is_break and j in (2, 3, 4)))
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=hair, left=hair, right=hair)
            if is_break:
                c.fill = PatternFill("solid", fgColor="D6EEF5")
            elif row["section"] != current_section and j == 2:
                c.font = Font(name="Montserrat", size=9, color=BLUE, bold=True)
        # row height from the longest cell (~ chars per line at the column width)
        # Montserrat 9pt fits roughly one character per unit of column width
        longest = max((len(str(v)) / max(COLUMNS[j][1] - 2, 4) for j, v in enumerate(vals)), default=1)
        ws.row_dimensions[r].height = max(18, min(170, 12.5 * (int(longest) + 1)))
        current_section = row["section"]
        r += 1

    ws.freeze_panes = ws.cell(row=head_row + 1, column=5)
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(COLUMNS))}{r - 1}"
    ws.sheet_view.zoomScale = 90

    # second sheet: storyline in reading order (section → titles), the "dash" view
    ws2 = wb.create_sheet("Storyline")
    ws2["A1"] = "Storyline — read the action titles top to bottom; they must argue the case on their own"
    ws2["A1"].font = Font(name="Montserrat", size=11, bold=True, color=NAVY)
    ws2.column_dimensions["A"].width = 6
    ws2.column_dimensions["B"].width = 18
    ws2.column_dimensions["C"].width = 100
    rr = 3
    for row in plan["slides"]:
        if row["type"] in ("cover", "agenda", "appendix"):
            continue
        ws2.cell(row=rr, column=1, value=row["n"]).font = Font(name="Montserrat", size=9, color=GREY)
        ws2.cell(row=rr, column=2, value=row["section"]).font = Font(name="Montserrat", size=9, color=BLUE, bold=True)
        c = ws2.cell(row=rr, column=3, value=row["title"])
        c.font = Font(name="Montserrat", size=10, color=INK, bold=row["type"] in ("divider", "exec_summary"))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        rr += 1

    # third sheet: data request derived from the plan
    ws3 = wb.create_sheet("Data request")
    heads = ("#", "Slide", "Data & sources", "Owner", "Status")
    for j, h in enumerate(heads, 1):
        c = ws3.cell(row=1, column=j, value=h.upper())
        c.font = Font(name="Montserrat", size=9, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
    for j, w in enumerate((6, 50, 60, 14, 12), 1):
        ws3.column_dimensions[get_column_letter(j)].width = w
    rr = 2
    for row in plan["slides"]:
        if not row["data"]:
            continue
        for j, v in enumerate((row["n"], row["title"], row["data"], row["owner"], "open"), 1):
            c = ws3.cell(row=rr, column=j, value=v)
            c.font = Font(name="Montserrat", size=9, color=INK)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        rr += 1

    wb.save(path)
    return path


# --------------------------------------------------------------------------- #
# Skeleton deck
# --------------------------------------------------------------------------- #
def build_skeleton(plan, path, draft=True):
    from aqmen_deck import Deck, Bullet, Palette, Font, MARGIN, CONTENT_TOP, CONTENT_BOTTOM, RAIL_LINE_X, SLIDE_W
    from pptx.util import Inches, Pt, Emu
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    d = Deck(draft=draft)
    sections = []
    for row in plan["slides"]:
        if row["type"] == "divider" and row["section"] not in sections:
            sections.append(row["section"])
    if not sections:
        for row in plan["slides"]:
            if row["type"] == "content" and row["section"] and row["section"] not in sections:
                sections.append(row["section"])

    def notes(slide, row):
        parts = []
        if row["purpose"]:
            parts.append("PURPOSE: " + row["purpose"])
        if row["exhibit"]:
            parts.append("EXHIBIT: " + row["exhibit"])
        if row["data"]:
            parts.append("DATA: " + row["data"])
        if row["owner"]:
            parts.append("OWNER: " + row["owner"])
        if row["image_prompt"]:
            parts.append("IMAGE PROMPT: " + row["image_prompt"])
        if parts:
            slide.notes_slide.notes_text_frame.text = "\n\n".join(parts)

    for row in plan["slides"]:
        t = row["type"]
        if t == "cover":
            s = d.title_slide(row["title"] or plan["title"], row.get("content") or plan.get("subtitle"), plan.get("date"))
        elif t == "agenda":
            s = d.agenda(sections + ["Appendix"])
        elif t == "divider":
            s = d.agenda(sections + ["Appendix"], active=row["section"])
        elif t == "exec_summary":
            rows = [(sec, [f"[{sec} — key messages to be written from the slides]"]) for sec in sections]
            s = d.executive_summary(rows, bottom_line=row["title"] if row["title"] else None)
        elif t == "appendix":
            s = d.agenda(sections + ["Appendix"], active="Appendix")
        else:
            s = d.content_slide(row["title"] or "[Action title]", body=None,
                                takeaways=[row["purpose"]] if row["purpose"] else None,
                                eyebrow=(row["section"], row.get("exhibit_short") or "") if row["section"] else None,
                                source=row["data"][:120] if row["data"] else None,
                                left_title=row.get("exhibit_title") or (row["exhibit"].split(":")[0][:70] if row["exhibit"] else "Exhibit"))
            # placeholder box describing the analysis to build
            x, y = MARGIN, CONTENT_TOP
            w = (RAIL_LINE_X - Inches(0.3) - MARGIN) if row["purpose"] else (SLIDE_W - 2 * MARGIN)
            h = CONTENT_BOTTOM - y
            box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
            box.fill.solid(); box.fill.fore_color.rgb = Palette.GREY_BG
            box.line.color.rgb = Palette.GREY_LINE; box.line.width = Pt(0.75)
            ln = box.line._get_or_add_ln()
            from lxml import etree
            from pptx.oxml.ns import qn
            etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
            box.shadow.inherit = False
            st = box._element.find(qn("p:style"))
            if st is not None:
                box._element.remove(st)
            lines = []
            if row["exhibit"]:
                lines.append([("EXHIBIT  ", {"bold": True, "color": Palette.NAVY}), (row["exhibit"], {})])
            if row["content"]:
                lines.append([("CONTENT  ", {"bold": True, "color": Palette.NAVY}), (row["content"], {})])
            if row["data"]:
                lines.append([("DATA  ", {"bold": True, "color": Palette.NAVY}), (row["data"], {})])
            if row["image_prompt"]:
                lines.append([("IMAGE PROMPT  ", {"bold": True, "color": Palette.BLUE}), (row["image_prompt"], {"italic": True})])
            d._text(s, Emu(x + Inches(0.25)), Emu(y + Inches(0.25)), Emu(w - Inches(0.5)), Emu(h - Inches(0.5)),
                    lines or "[analysis to be built]", size=10, color=Palette.INK, line_spacing=1.05,
                    space_after=8, autofit=False)
        notes(s, row)
    d.save(path)
    return len(d.prs.slides._sldIdLst)


# --------------------------------------------------------------------------- #
# Derive a plan from a cdd-output content file
# --------------------------------------------------------------------------- #
KIND_LABEL = {
    "content": "Text slide", "scenario": "Scenario slide", "table": "Table", "chart": "Chart",
    "waterfall": "Waterfall", "mekko": "Marimekko", "driver_tree": "Driver tree",
    "positioning": "Archetype 2×2", "harvey": "Harvey-ball matrix", "heatmap": "Sensitivity heatmap",
    "revenue_build": "Revenue build",
}


def from_cdd(content_path, out_path):
    c = json.loads(Path(content_path).read_text(encoding="utf-8"))
    slides = [
        {"type": "cover", "section": "", "title": c["title"], "content": c.get("subtitle", "")},
        {"type": "agenda", "section": "", "title": "Agenda"},
        {"type": "exec_summary", "section": "", "title": c.get("bottom_line", "Executive summary"),
         "purpose": "State the answer and the decision it informs before any exhibit"},
    ]
    for part in c["parts"]:
        slides.append({"type": "divider", "section": part["name"], "title": part["name"]})
        for sec in part["sections"]:
            kind = sec["kind"]
            exhibit = KIND_LABEL.get(kind, kind)
            detail = sec.get("chart_title") or sec.get("left_title") or sec["title"]
            body = sec.get("body") or []
            content = "; ".join((b[0] if isinstance(b, list) else str(b)) for b in body[:3])
            slides.append({
                "type": "content", "section": part["name"], "title": sec["headline"],
                "content": content or detail, "exhibit": f"{exhibit}: {detail}",
                "exhibit_title": detail,
                "data": sec.get("source", ""), "owner": "agent" if kind in ("chart", "mekko", "driver_tree", "table") else "analyst",
                "purpose": " · ".join(sec.get("so_whats", [])[:2]),
                "image_prompt": "",
            })
    slides.append({"type": "appendix", "section": "Appendix", "title": "Sources & confidence"})
    plan = {"title": c["title"], "subtitle": "Presentation plan (Dot-Dash) derived from the content file",
            "client": c.get("client", ""), "date": c.get("date", ""), "slug": c.get("slug", ""), "slides": slides}
    Path(out_path).write_text(json.dumps(plan, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"wrote {out_path} ({len(slides)} rows)")


def main(argv):
    if "--from-cdd" in argv:
        i = argv.index("--from-cdd")
        return from_cdd(argv[i + 1], argv[i + 2])
    if len(argv) < 3:
        raise SystemExit(__doc__)
    plan = load_plan(argv[1])
    outdir = Path(argv[2]); outdir.mkdir(parents=True, exist_ok=True)
    only = argv[argv.index("--only") + 1] if "--only" in argv else None
    final = "--final" in argv
    base = plan["slug"]
    if only in (None, "xlsx"):
        p = outdir / f"{base} - Dot-Dash.xlsx"
        build_xlsx(plan, str(p)); print(f"xlsx  {p}")
    if only in (None, "deck"):
        p = outdir / f"{base} - Skeleton deck.pptx"
        n = build_skeleton(plan, str(p), draft=not final); print(f"deck  {p}  ({n} slides)")


if __name__ == "__main__":
    main(sys.argv)
