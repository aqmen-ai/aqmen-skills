"""
Render the aqmen CDD deliverable from one content JSON into three formats:

    deck    -> editable .pptx via aqmen_deck.py (native charts, house style)
    html    -> single self-contained .html report (ECharts, house style)
    summary -> 2-page executive summary .docx (Montserrat, brand palette)

One content file, three outputs, so the deck, the report and the leave-behind
never disagree. The content schema is documented in ../references/content-format.md
and exemplified in ../assets/example-carwash.json.

The deck and HTML builders are ports of the plugin's own template generators
(scripts/build-templates.py and build-html-examples.py), adapted to (a) read
JSON, (b) group sections into the CDD parts (Context / Market / Competition /
Company) with an agenda slide per part, and (c) add the table and waterfall
kinds the storyline needs.
"""
from __future__ import annotations

import html as _html
import json
import math
import os
import re
import sys
from datetime import date
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
REFS = SKILL / "references"
ASSETS = SKILL / "assets"
sys.path.insert(0, str(REFS))

from aqmen_deck import (  # noqa: E402
    Deck, Bullet, ChartSpec, MekkoColumn, HarveyRow, DriverNode, DriverColumn, MatrixItem,
)

YEAR = str(date.today().year)
EXHIBIT_KINDS = {"chart", "mekko", "harvey", "driver_tree", "positioning", "heatmap",
                 "revenue_build", "table", "waterfall"}


# --------------------------------------------------------------------------- #
# Content helpers
# --------------------------------------------------------------------------- #
# The fields that carry a section's data, per kind: a section with a `ref` must
# have them filled (resolved from the workspace) before it can render.
DATA_FIELDS = {
    "content": ("body",), "scenario": ("body",), "table": ("columns", "rows"),
    "chart": ("chart",), "waterfall": ("steps",), "mekko": ("mekko",),
    "driver_tree": ("tree",), "positioning": ("items",), "harvey": ("columns", "rows"),
    "heatmap": ("cols", "rows", "values"), "revenue_build": ("groups",),
}


def _check_ref(part, sec):
    """A `ref` points at a workspace chart or spreadsheet range; the agent
    resolves it into literal values before building. Refuse unresolved refs
    rather than render an empty exhibit."""
    ref = sec.get("ref")
    if not ref:
        return
    missing = [f for f in DATA_FIELDS.get(sec.get("kind"), ()) if not sec.get(f)]
    if missing or not ref.get("resolved"):
        what = ", ".join(missing + ([] if ref.get("resolved") else ["ref.resolved"]))
        raise SystemExit(
            f"section {sec.get('title')!r} in part {part['name']!r} has a ref "
            f"{json.dumps({k: v for k, v in ref.items() if k != 'resolved'})} but no resolved data "
            f"(missing: {what}). Resolve it with get_chart / read_spreadsheet, write the values "
            f"and ref.resolved into the content, then build.")


def load_content(path):
    c = json.loads(Path(path).read_text(encoding="utf-8"))
    c.setdefault("doctype", "Commercial Due Diligence")
    c.setdefault("date", YEAR)
    c.setdefault("confidence", 4)
    c.setdefault("kpis", [])
    c.setdefault("sources", [])
    for part in c["parts"]:
        part.setdefault("sections", [])
        for sec in part["sections"]:
            sec.setdefault("so_whats", [])
            sec.setdefault("body", [])
            _check_ref(part, sec)
    return c


def _is_appendix(part):
    return part["name"].strip().lower() == "appendix"


def split_appendix(content):
    """(parts, appendix): the storyline parts, and the built-in Appendix with
    any user part named "Appendix" (any case) merged in. Its sections render
    after the Appendix divider and before Sources & confidence."""
    parts = [p for p in content["parts"] if not _is_appendix(p)]
    extra = [p for p in content["parts"] if _is_appendix(p)]
    appendix = {"name": "Appendix",
                "subitems": [i for p in extra for i in p.get("subitems", [])],
                "intro": " ".join(p["intro"] for p in extra if p.get("intro")),
                "sections": [s for p in extra for s in p.get("sections", [])]}
    return parts, appendix


def _body_items(body):
    """Normalise body items to (text, bold, level). Accepts strings or lists."""
    out = []
    for it in body or []:
        if isinstance(it, str):
            out.append((it, False, 0))
        else:
            text = it[0]
            bold = bool(it[1]) if len(it) > 1 else False
            level = int(it[2]) if len(it) > 2 else 0
            out.append((text, bold, level))
    return out


def _takeaways(sec):
    tk = list(sec.get("so_whats", []))
    if sec.get("watchout"):
        tk.append("Watch-out: " + sec["watchout"])
    return tk


def all_sections(content):
    parts, appendix = split_appendix(content)
    for part in parts + [appendix]:
        for sec in part["sections"]:
            yield part, sec


def slug(content):
    base = content.get("slug") or content["title"]
    return re.sub(r"[^A-Za-z0-9]+", "-", base).strip("-")


# --------------------------------------------------------------------------- #
# Deck
# --------------------------------------------------------------------------- #
def _waterfall_series(steps):
    """Turn [(label, delta_or_total, is_total)] into a stacked column with an
    invisible base, so the bridge renders as a waterfall with native charts."""
    cats, base, up, down, tot = [], [], [], [], []
    running = 0.0
    for lbl, val, is_total in steps:
        cats.append(lbl)
        if is_total:
            running = val
            base.append(0); up.append(0); down.append(0); tot.append(val)
        else:
            if val >= 0:
                base.append(running); up.append(val); down.append(0); tot.append(0)
                running += val
            else:
                running += val
                base.append(running); up.append(0); down.append(-val); tot.append(0)
    return cats, [("Total", tot), ("Increase", up), ("Decrease", down), ("_base", base)]


def table_slide(d, headline, columns, rows, takeaways=None, eyebrow=None, source=None,
                illustrative=False, chart_title=None, widths=None):
    """A native PowerPoint table in the house style (navy header row, thin navy
    rules, Montserrat), laid out with the same chrome, headline and takeaways
    rail as the builder's other slides. Used for trends, KPC and option tables.
    Lives here until the shared builder grows a table method."""
    import aqmen_deck as AD
    from pptx.util import Inches, Pt, Emu
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.dml.color import RGBColor

    d._page += 1
    slide = d._slide()
    d._chrome(slide, eyebrow=eyebrow, source=source)
    d._headline(slide, headline, illustrative=illustrative)
    x = AD.MARGIN
    y = AD.CONTENT_TOP
    w = (AD.RAIL_LINE_X - Inches(0.3) - AD.MARGIN) if takeaways else (AD.SLIDE_W - 2 * AD.MARGIN)
    h = AD.CONTENT_BOTTOM - y
    d._header_row(slide, chart_title, rail=bool(takeaways))
    n_rows, n_cols = len(rows) + 1, len(columns)
    shape = slide.shapes.add_table(n_rows, n_cols, x, y, w, Inches(0.4) * n_rows)
    tbl = shape.table
    # widths: explicit weights, else proportional to the text each column carries
    # (sqrt-damped so a long rationale column doesn't starve short ones).
    if not widths:
        import math
        lens = []
        for j in range(n_cols):
            cells = [str(columns[j])] + [str(r[j]) if j < len(r) else "" for r in rows]
            lens.append(max(6.0, sum(len(c) for c in cells) / len(cells)))
        widths = [math.sqrt(v) for v in lens]
    total = sum(widths)
    for j, wt in enumerate(widths):
        tbl.columns[j].width = Emu(int(w * wt / total))
    # minimal row heights; text grows them. Fit check: ~10 body rows max.
    for i in range(n_rows):
        tbl.rows[i].height = Inches(0.32)
    tbl.first_row = True
    # turn off the default banded style so our fills show
    tblPr = shape._element.graphic.graphicData.tbl.tblPr
    tblPr.set("bandRow", "0")
    navy = RGBColor.from_string("03045E")
    ink = RGBColor.from_string("302E2E")
    white = RGBColor.from_string("FFFFFF")
    for j, name in enumerate(columns):
        c = tbl.cell(0, j)
        c.fill.solid(); c.fill.fore_color.rgb = navy
        c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.04)
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = str(name).upper()
        r.font.size = Pt(9.5); r.font.bold = True; r.font.color.rgb = white; r.font.name = AD.Font.HEAD
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
    body_size = 10 if len(rows) <= 4 else (9 if len(rows) <= 7 else 8)
    for i, row in enumerate(rows, start=1):
        for j in range(n_cols):
            c = tbl.cell(i, j)
            c.fill.solid(); c.fill.fore_color.rgb = white if i % 2 else RGBColor.from_string("F4F9FC")
            c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.03)
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run(); r.text = str(row[j]) if j < len(row) else ""
            r.font.size = Pt(body_size); r.font.color.rgb = ink; r.font.name = AD.Font.BODY
            r.font.bold = (j == 0)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
    if takeaways:
        d._rail(slide, takeaways)
    return slide


def _hide_waterfall_base(slide):
    """Make the first ('_base') series of the slide's chart invisible: no fill,
    no line, no data labels, and drop it from the legend."""
    from pptx.util import Pt as _Pt  # noqa: F401
    for shape in slide.shapes:
        if not getattr(shape, "has_chart", False) or not shape.has_chart:
            continue
        chart = shape.chart
        plot = chart.plots[0]
        base = plot.series[0]
        base.format.fill.background()
        base.format.line.fill.background()
        # data labels: turn off on the base series only
        dLbls = base._element.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/chart}dLbls")
        if dLbls is not None:
            base._element.remove(dLbls)
        # legend: delete the base entry (idx 0)
        if chart.has_legend:
            legend = chart.legend._element
            from pptx.oxml import parse_xml
            ns = "http://schemas.openxmlformats.org/drawingml/2006/chart"
            entry = parse_xml(f'<c:legendEntry xmlns:c="{ns}"><c:idx val="0"/><c:delete val="1"/></c:legendEntry>')
            pos = legend.find(f"{{{ns}}}legendPos")
            (pos.addnext(entry) if pos is not None else legend.insert(0, entry))
        break


def build_deck(content, path, draft=True):
    d = Deck(draft=draft)
    doctype = content["doctype"]
    d.title_slide(content["title"], content.get("subtitle"), content.get("date", YEAR))
    parts, appendix = split_appendix(content)
    part_names = [p["name"] for p in parts] + ["Appendix"]
    subitems = {p["name"]: p["subitems"] for p in parts + [appendix] if p.get("subitems")}
    d.agenda(part_names, subitems=subitems)
    d.executive_summary([(lbl, list(bul)) for lbl, bul in content["exec_summary"]],
                        bottom_line=content.get("bottom_line"))

    for part in parts + [appendix]:
        d.agenda(part_names, active=part["name"], subitems=subitems)
        for sec in part["sections"]:
            kind = sec["kind"]
            eyebrow = (part["name"], sec["title"])
            common = dict(eyebrow=eyebrow, source=sec.get("source"),
                          takeaways=_takeaways(sec) or None,
                          illustrative=sec.get("illustrative", False))
            body = [Bullet(t, level=l, bold=b) for (t, b, l) in _body_items(sec.get("body"))]
            exhibit = sec.get("chart_title") or sec.get("left_title") or sec["title"]
            if kind in ("content", "scenario"):
                d.content_slide(sec["headline"], body=body,
                                left_title=sec.get("left_title") or sec["title"], **common)
            elif kind == "table":
                table_slide(d, sec["headline"], sec["columns"], sec["rows"],
                            chart_title=exhibit, widths=sec.get("widths"), **common)
            elif kind == "chart":
                c = sec["chart"]
                spec = ChartSpec(kind=c["kind"], categories=list(c["categories"]),
                                 series=[(n, list(v)) for n, v in c["series"]],
                                 number_format=c.get("number_format", "#,##0"),
                                 legend=c.get("legend", True), y_title=c.get("y_title"))
                # CAGR annotations on year-based column charts unless switched off
                cagr = sec.get("cagr", c["kind"] in ("column", "stacked_column"))
                d.chart_slide(sec["headline"], chart=spec, chart_title=exhibit, cagr=cagr, **common)
            elif kind == "waterfall":
                steps = [(s[0], float(s[1]), bool(s[2]) if len(s) > 2 else False) for s in sec["steps"]]
                cats, series = _waterfall_series(steps)
                # Base series first so the deltas float; it is then made invisible.
                ordered = [series[3], series[0], series[1], series[2]]
                spec = ChartSpec(kind="stacked_column", categories=cats, series=ordered,
                                 number_format=sec.get("number_format", "#,##0"), legend=True)
                d.chart_slide(sec["headline"], chart=spec, chart_title=exhibit, **common)
                _hide_waterfall_base(d.prs.slides[-1])
            elif kind == "mekko":
                cols = [MekkoColumn(lbl, w, [tuple(s) for s in segs]) for (lbl, w, segs) in sec["mekko"]]
                d.marimekko_slide(sec["headline"], columns=cols, chart_title=exhibit, **common)
            elif kind == "harvey":
                rows = [HarveyRow(lbl, list(cells)) for (lbl, cells) in sec["rows"]]
                d.harvey_matrix_slide(sec["headline"], columns=sec["columns"], rows=rows,
                                      row_label_header=sec.get("row_header", ""), **common)
                d._header_row(d.prs.slides[-1], exhibit, rail=bool(common["takeaways"]))
            elif kind == "driver_tree":
                cols = [DriverColumn(c["header"],
                                     [DriverNode(n["label"], n.get("value"), n.get("certainty"),
                                                 n.get("expr"), n.get("parent", 0), n.get("operator"),
                                                 n.get("short_value"))
                                      for n in c["nodes"]]) for c in sec["tree"]]
                d.driver_tree_slide(sec["headline"], columns=cols, chart_title=exhibit,
                                    note=sec.get("note"), **common)
            elif kind == "positioning":
                items = [MatrixItem(i["label"], i.get("x", 0.5), i.get("y", 0.5), i.get("size", 0.5),
                                    i.get("num"),
                                    [tuple(p) for p in i["points"]] if i.get("points") else None,
                                    i.get("pad", 0.05), i.get("show_points", True))
                         for i in sec["items"]]
                d.positioning_matrix_slide(sec["headline"], items=items, x_title=sec["x_title"],
                                           y_title=sec["y_title"], x_ticks=sec.get("x_ticks"),
                                           y_ticks=sec.get("y_ticks"),
                                           chart_title=exhibit, **common)
            elif kind == "heatmap":
                d.heatmap_slide(sec["headline"], cols=sec["cols"], rows=sec["rows"],
                                values=sec["values"], base=tuple(sec["base"]) if sec.get("base") else None,
                                col_header=sec.get("col_header", ""), row_header=sec.get("row_header", ""),
                                value_fmt=sec.get("value_fmt", "{:.0f}"),
                                chart_title=exhibit, **common)
            elif kind == "revenue_build":
                common_nt = {k: v for k, v in common.items() if k != "takeaways"}
                d.revenue_build_slide(
                    sec["headline"],
                    groups=[(g["archetype"], g["players"]) for g in sec["groups"]],
                    info_cols=[tuple(c) for c in sec.get("info_cols", [("hq", "HQ"), ("employees", "Employees")])],
                    total=sec.get("total"),
                    total_label=sec.get("total_label", "Total addressable (bottom-up market)"),
                    unit=sec.get("unit", "$m"), chart_title=exhibit, **common_nt)
            else:
                raise ValueError(f"unknown section kind: {kind!r} ({sec.get('title')})")

    bullets = [Bullet(f"{src}  —  confidence {conf}/5  —  {note}") for (src, conf, note) in content["sources"]]
    d.content_slide("Sources & confidence", body=bullets, eyebrow=("Appendix", "Sources & confidence"),
                    source="Confidence: 5 primary · 4 credible secondary · 3 triangulated (news cap) · 2 weak · 1 estimate")
    d.save(path)
    return len(d.prs.slides._sldIdLst)


# --------------------------------------------------------------------------- #
# HTML
# --------------------------------------------------------------------------- #
def esc(s):
    return _html.escape(str(s))


def shell_parts(title):
    txt = (ASSETS / "report-shell.html").read_text(encoding="utf-8")
    head = txt[: txt.index("</head>") + len("</head>")]
    head = re.sub(r"<!--.*?-->\s*", "", head, count=1, flags=re.S)
    head = head.replace("{{REPORT_TITLE}}", esc(title))
    open_idx = txt.rindex("<script>") + len("<script>")
    script = txt[open_idx: txt.rindex("</script>")]
    helpers = script[: script.index("// Example")].rstrip()
    scenario_js = script[script.index("// --- Scenario tabs ---"):].rstrip()
    return head, helpers, scenario_js


def _series_bar(series, stacked=False, horizontal=False):
    out = []
    for name, vals in series:
        s = {"name": name, "type": "bar", "data": list(vals),
             "label": {"show": True, "fontSize": 10, "position": "right" if horizontal else "top"}}
        if stacked:
            s["stack"] = "total"
            s["label"]["position"] = "inside"
        out.append(s)
    return out


def chart_option(chart):
    kind = chart["kind"]
    cats = list(chart["categories"])
    series = chart["series"]
    multi = len(series) > 1
    legend = chart.get("legend", True) and multi
    opt = {"tooltip": {"trigger": "axis", "axisPointer": {"type": "shadow"}},
           "grid": {"left": 48, "right": 24, "top": 20, "bottom": 48 if legend else 28}}
    if legend:
        opt["legend"] = {"bottom": 0}
    if kind == "line":
        opt["xAxis"] = {"type": "category", "data": cats}
        opt["yAxis"] = {"type": "value"}
        opt["series"] = [{"name": n, "type": "line", "smooth": False, "data": list(v),
                          "label": {"show": True, "fontSize": 10, "position": "top"}} for n, v in series]
    elif kind == "bar":
        opt["yAxis"] = {"type": "category", "data": cats}
        opt["xAxis"] = {"type": "value"}
        opt["series"] = _series_bar(series, stacked=False, horizontal=True)
    else:
        stacked = kind in ("stacked_column", "stacked_bar")
        opt["xAxis"] = {"type": "category", "data": cats}
        opt["yAxis"] = {"type": "value"}
        opt["series"] = _series_bar(series, stacked=stacked, horizontal=False)
    return opt


def waterfall_option(sec):
    steps = [(s[0], float(s[1]), bool(s[2]) if len(s) > 2 else False) for s in sec["steps"]]
    cats, series = _waterfall_series(steps)
    d = dict(series)
    return {"tooltip": {"trigger": "axis", "axisPointer": {"type": "shadow"}},
            "grid": {"left": 48, "right": 24, "top": 20, "bottom": 28},
            "xAxis": {"type": "category", "data": cats},
            "yAxis": {"type": "value"},
            "series": [
                {"name": "base", "type": "bar", "stack": "w", "itemStyle": {"color": "transparent"},
                 "emphasis": {"itemStyle": {"color": "transparent"}}, "data": d["_base"]},
                {"name": "Total", "type": "bar", "stack": "w", "data": d["Total"],
                 "itemStyle": {"color": "#03045E"}, "label": {"show": True, "position": "top", "fontSize": 10}},
                {"name": "Increase", "type": "bar", "stack": "w", "data": d["Increase"],
                 "itemStyle": {"color": "#2E6FD6"}, "label": {"show": True, "position": "top", "fontSize": 10}},
                {"name": "Decrease", "type": "bar", "stack": "w", "data": d["Decrease"],
                 "itemStyle": {"color": "#ADE8F3"}, "label": {"show": True, "position": "top", "fontSize": 10}},
            ]}


def _ball(fill):
    f = max(0.0, min(1.0, fill))
    nv = "#03045E"
    ring = f'<circle cx="8" cy="8" r="6.6" fill="#fff" stroke="{nv}" stroke-width="1.3"/>'
    if f <= 0:
        wedge = ""
    elif f >= 1:
        wedge = f'<circle cx="8" cy="8" r="6.6" fill="{nv}"/>'
    else:
        ang = 2 * math.pi * f
        ex, ey = 8 + 6.6 * math.sin(ang), 8 - 6.6 * math.cos(ang)
        large = 1 if f > 0.5 else 0
        wedge = f'<path d="M8,8 L8,1.4 A6.6,6.6 0 {large} 1 {ex:.2f},{ey:.2f} Z" fill="{nv}"/>'
    return f'<svg width="15" height="15" viewBox="0 0 16 16" style="vertical-align:middle">{ring}{wedge}</svg>'


def _bullets_html(body):
    lis = []
    for (text, bold, level) in _body_items(body):
        inner = f"<strong>{esc(text)}</strong>" if bold else esc(text)
        ml = f' style="margin-left:{level*18}px"' if level else ""
        lis.append(f"<li{ml}>{inner}</li>")
    return "<ul>\n" + "\n".join(lis) + "\n</ul>"


def _insight_html(so_whats, watchout):
    out = []
    if so_whats:
        out.append('<div class="callout insight"><span class="tag">Insight</span>' + esc(" · ".join(so_whats)) + "</div>")
    if watchout:
        out.append('<div class="callout watchout"><span class="tag">⚠ Watch-out</span>' + esc(watchout) + "</div>")
    return "\n".join(out)


def _source_line(source):
    return (f'<p style="font-size:12.5px;color:var(--muted);margin:6px 0 0">Source: {esc(source)}</p>') if source else ""


def _table_html(sec):
    head = "<tr>" + "".join(f"<th>{esc(c)}</th>" for c in sec["columns"]) + "</tr>"
    body = []
    for r in sec["rows"]:
        tds = "".join((f"<td><b>{esc(c)}</b></td>" if j == 0 else f"<td>{esc(c)}</td>") for j, c in enumerate(r))
        body.append(f"<tr>{tds}</tr>")
    return "<table><thead>" + head + "</thead><tbody>" + "".join(body) + "</tbody></table>"


def _harvey_table(sec):
    cols = sec["columns"]
    head = ("<tr><th>" + esc(sec.get("row_header", "")) + "</th>" + "".join(f"<th>{esc(c)}</th>" for c in cols) + "</tr>")
    body = []
    for (label, cells) in sec["rows"]:
        tds = "".join(f'<td style="text-align:center">{_ball(v)}</td>' for v in cells)
        body.append(f"<tr><td><b>{esc(label)}</b></td>{tds}</tr>")
    return "<table><thead>" + head + "</thead><tbody>" + "".join(body) + "</tbody></table>"


def _driver_tree_html(sec):
    parts = ['<div class="dtree">']
    for col in sec["tree"]:
        parts.append('<div class="dtcol">')
        parts.append(f'<div class="dthead">{esc(col["header"])}</div>')
        for node in col["nodes"]:
            cert = node.get("certainty")
            cdot = f'<span class="cert cert-{cert}"></span>' if cert is not None else ""
            expr = f'<div class="expr">{esc(node["expr"])}</div>' if node.get("expr") else ""
            val = f'<div class="val">{esc(node["value"])}</div>' if node.get("value") is not None else ""
            op = f'<span class="op">{esc(node["operator"])}</span>' if node.get("operator") else ""
            parts.append(f'<div class="dtnode">{cdot}{op}<div class="lab">{esc(node["label"])}</div>{expr}{val}</div>')
        parts.append("</div>")
    parts.append("</div>")
    parts.append('<div class="dtkey"><span>Certainty of assumptions:</span>'
                 '<span><i class="cert-0"></i>Low</span><span><i class="cert-1"></i>Medium</span>'
                 '<span><i class="cert-2"></i>High</span></div>')
    if sec.get("note"):
        parts.append(f'<div class="dtnote">▸ {esc(sec["note"])}</div>')
    return "\n".join(parts)


def _positioning_html(sec):
    items = sec["items"]
    xt = sec.get("x_ticks", []) or []
    yt = sec.get("y_ticks", []) or []
    cells = []
    for it in items:
        badge = f'<span class="badge">{esc(it["num"])}</span>' if it.get("num") else ""
        pts = it.get("points") or []
        pad = it.get("pad", 0.05)
        if pts:
            minx = max(0.0, min(p[0] for p in pts) - pad); maxx = min(1.0, max(p[0] for p in pts) + pad)
            miny = max(0.0, min(p[1] for p in pts) - pad); maxy = min(1.0, max(p[1] for p in pts) + pad)
        else:
            hw = 0.5 * it.get("size", 0.5) * 0.5
            minx, maxx = it.get("x", 0.5) - hw, it.get("x", 0.5) + hw
            miny, maxy = it.get("y", 0.5) - hw * 1.2, it.get("y", 0.5) + hw * 1.2
        cells.append(f'<div class="pmitem" style="left:{minx*100:.1f}%;bottom:{miny*100:.1f}%;'
                     f'width:{(maxx-minx)*100:.1f}%;height:{(maxy-miny)*100:.1f}%">{badge}<div class="arch">{esc(it["label"])}</div></div>')
        if it.get("show_points", True):
            for (fx, fy) in pts:
                cells.append(f'<div class="pmdot" style="left:{fx*100:.1f}%;bottom:{fy*100:.1f}%"></div>')
    for i, tk in enumerate(xt):
        if tk:
            cells.append(f'<div class="pmxtick" style="left:{(i+0.5)/len(xt)*100:.1f}%">{esc(tk)}</div>')
    for i, tk in enumerate(yt):
        if tk:
            cells.append(f'<div class="pmytick" style="bottom:{(i+0.5)/len(yt)*100:.1f}%">{esc(tk)}</div>')
    return (f'<div class="pmatrix"><div class="pmytitle">{esc(sec["y_title"])}</div>'
            f'<div class="pmarea">{"".join(cells)}</div></div><div class="pmxtitle">{esc(sec["x_title"])}</div>')


_MK_PALETTE = ["#03045E", "#0728A3", "#2E6FD6", "#6BAED6", "#ADE8F3", "#8A8A8A"]


def _mekko_html(sec):
    mekko = sec["mekko"]
    seg_names = [s for s, _ in mekko[0][2]]
    color, ci = {}, 0
    for s in seg_names:
        if "whitespace" in s.lower():
            color[s] = None
        else:
            color[s] = _MK_PALETTE[ci % len(_MK_PALETTE)]; ci += 1
    total_w = sum(w for _, w, _ in mekko) or 1.0
    max_total = max(sum(v for _, v in segs) for _, _, segs in mekko) or 1.0
    cols = []
    for (lab, w, segs) in mekko:
        ctot = sum(v for _, v in segs) or 1.0
        segdivs = []
        for (s, v) in segs:
            h = v / ctot * 100
            cls = "mkseg ws" if color[s] is None else "mkseg"
            bg = "" if color[s] is None else f";background:{color[s]}"
            segdivs.append(f'<div class="{cls}" style="height:{h:.1f}%{bg}">{v:g}</div>')
        colpct = sum(v for _, v in segs) / max_total * 100
        cols.append(f'<div class="mkcol" style="flex:{w:.4f} 1 0"><div class="mkhead">{sum(v for _, v in segs):g}'
                    f'<span>({w / total_w * 100:.0f}%)</span></div><div class="mkstackarea"><div class="mkstack" '
                    f'style="height:{colpct:.1f}%">{"".join(segdivs)}</div></div><div class="mkxlab">{esc(lab)}</div></div>')
    leg = []
    for s in seg_names:
        sw = '<i class="wslegend"></i>' if color[s] is None else f'<i style="background:{color[s]}"></i>'
        leg.append(f"<span>{sw}{esc(s)}</span>")
    return f'<div class="mekko">{"".join(cols)}</div><div class="mklegend">{"".join(leg)}</div>'


def _revenue_build_html(sec):
    info_cols = sec.get("info_cols", [["hq", "HQ"], ["employees", "Employees"]])
    unit = sec.get("unit", "$m")
    groups = sec["groups"]
    total = sec.get("total")
    fmt = lambda v: f"{v:,.0f}"  # noqa: E731
    ncols = len(info_cols) + 4
    all_m = [p.get("market_rev") or 0 for g in groups for p in g["players"]]
    mx = max(all_m) or 1.0

    def bar(v, full=False):
        w = 100.0 if full else ((v / mx * 100) if v is not None else 0)
        val = fmt(v) if v is not None else "—"
        return (f'<div class="rbar"><span class="rtrack"><span class="rfill" style="width:{min(w,100):.1f}%"></span></span><b>{val}</b></div>')

    heads = ('<th class="l">Company</th>' + "".join(f'<th class="l">{esc(h)}</th>' for _, h in info_cols)
             + f'<th>Revenue ({unit})</th><th>% to market</th><th class="l">Market revenue ({unit})</th>')
    rows = []
    for g in groups:
        rows.append(f'<tr class="grp"><td colspan="{ncols}">{esc(g["archetype"])}</td></tr>')
        for p in g["players"]:
            info_tds = "".join(f'<td>{esc(str(p.get(k, "")))}</td>' for k, _ in info_cols)
            rev, pct, m = p.get("revenue"), p.get("pct"), p.get("market_rev")
            rev_td = f'<td class="num">{fmt(rev) if rev is not None else "—"}</td>'
            pct_td = f'<td class="num">{pct:g}%</td>' if pct is not None else '<td class="num">—</td>'
            rows.append(f'<tr><td><b>{esc(p.get("name",""))}</b></td>{info_tds}{rev_td}{pct_td}<td>{bar(m)}</td></tr>')
    if total is not None:
        lbl = esc(sec.get("total_label", "Total addressable (bottom-up market)"))
        rows.append(f'<tr class="total"><td colspan="{ncols-1}">{lbl}</td><td>{bar(total, full=True)}</td></tr>')
    cols = '<colgroup>' + '<col>' * (ncols - 1) + '<col class="mrev"></colgroup>'
    return f'<table class="rbuild">{cols}<thead><tr>{heads}</tr></thead><tbody>{"".join(rows)}</tbody></table>'


def _heatmap_html(sec):
    rows, cols, vals = sec["rows"], sec["cols"], sec["values"]
    base = sec.get("base")
    fmt = sec.get("value_fmt", "{:.0f}")
    flat = [v for r in vals for v in r]
    lo, hi = min(flat), max(flat)
    span = (hi - lo) or 1.0

    def cellcol(t):
        a, b = (0xAD, 0xE8, 0xF3), (0x03, 0x04, 0x5E)
        return "#%02X%02X%02X" % tuple(int(a[k] + (b[k] - a[k]) * t) for k in range(3))

    head = (f'<tr><th class="rowlab">{esc(sec.get("row_header",""))}</th>' + "".join(f'<th class="collab">{esc(c)}</th>' for c in cols) + "</tr>")
    body = []
    for i, rlab in enumerate(rows):
        tds = [f'<td class="rowlab">{esc(rlab)}</td>']
        for j, c in enumerate(cols):
            v = vals[i][j]
            t = (v - lo) / span
            isb = base and tuple(base) == (i, j)
            ink = "#302E2E" if t <= 0.55 else "#fff"
            tds.append(f'<td class="cell{" base" if isb else ""}" style="background:{cellcol(t)};color:{ink}">{esc(fmt.format(v))}</td>')
        body.append("<tr>" + "".join(tds) + "</tr>")
    return (f'<div class="hmaxis col">{esc(sec.get("col_header",""))}</div><div class="hmwrap"><table class="heatmap"><thead>'
            + head + "</thead><tbody>" + "".join(body) + "</tbody></table></div>")


def build_html(content):
    head, helpers, scenario_js = shell_parts(content["title"])
    charts = []
    parts = ['<div class="topbar"></div>', '<div class="page">']
    parts.append('<div class="brandbar"><span class="logo" role="img" aria-label="Aqmen"></span>'
                 f'<span class="wordmark">Aqmen</span><span class="doctype">{esc(content["doctype"])}</span></div>')
    conf = content.get("confidence", 4)
    parts.append('<header class="cover">'
                 f'<h1>{esc(content["title"])}</h1><p class="subtitle">{esc(content.get("subtitle", ""))}</p><div class="meta">'
                 f'<span><b>Project:</b> {esc(content.get("project", content.get("client", "")))}</span>'
                 f'<span><b>Date:</b> {esc(content.get("date", YEAR))}</span>'
                 f'<span><b>Prepared by:</b> {esc(content.get("author", "Aqmen"))}</span>'
                 f'<span><b>Overall confidence:</b> <span class="cf cf-{conf}">{conf}</span></span></div></header>')
    parts.append(f'<section class="bottomline"><span class="lbl">Bottom line</span><p>{esc(content["bottom_line"])}</p></section>')
    groups = []
    for (label, bullets) in content["exec_summary"]:
        subs = "".join(f"<li>{esc(b)}</li>" for b in bullets)
        groups.append(f"<li><b>{esc(label)}</b><ul>{subs}</ul></li>")
    parts.append('<section class="exec"><h2>Executive summary</h2><ul>' + "".join(groups) + "</ul></section>")
    tiles = []
    for (val, label, delta) in content.get("kpis", []):
        d = f'<div class="delta up">▲ {esc(str(delta).lstrip("▲ "))}</div>' if delta else ""
        tiles.append(f'<div class="kpi"><div class="val">{esc(val)}</div><div class="lbl">{esc(label)}</div>{d}</div>')
    if tiles:
        parts.append('<div class="kpis">' + "".join(tiles) + "</div>")

    story, appendix = split_appendix(content)
    chapters = story + ([appendix] if appendix["sections"] else [])
    for pi, part in enumerate(chapters, start=1):
        parts.append(f'<h2 style="margin-top:36px;padding-top:18px;border-top:2px solid var(--brand)">'
                     f'<span class="num">{pi}.</span>{esc(part["name"])}</h2>')
        if part.get("intro"):
            parts.append(f"<p>{esc(part['intro'])}</p>")
        for si, sec in enumerate(part["sections"], start=1):
            kind = sec["kind"]
            parts.append("<section>")
            parts.append(f'<h3 style="color:var(--brand);margin:22px 0 6px">{pi}.{si} {esc(sec["title"])}</h3>')
            parts.append(f"<p><strong>{esc(sec['headline'])}</strong></p>")
            if sec.get("left_title"):
                parts.append(f'<p style="font-weight:600;color:var(--brand);margin:12px 0 4px">{esc(sec["left_title"])}</p>')
            if sec.get("body"):
                parts.append(_bullets_html(sec["body"]))
            cap = sec.get("chart_title", "")
            illus = " · ILLUSTRATIVE" if sec.get("illustrative") else ""
            caption = f'<figcaption>{esc(cap)} · Source: {esc(sec.get("source",""))}{illus}</figcaption>'
            if kind in ("chart", "waterfall"):
                cid = f"chart{len(charts) + 1}"
                charts.append((cid, chart_option(sec["chart"]) if kind == "chart" else waterfall_option(sec)))
                parts.append(f'<figure><div id="{cid}" class="chart"></div>{caption}</figure>')
            elif kind == "mekko":
                parts.append(_mekko_html(sec)); parts.append(caption)
            elif kind == "harvey":
                parts.append(_harvey_table(sec)); parts.append(_source_line(sec.get("source", "")))
            elif kind == "table":
                parts.append(_table_html(sec)); parts.append(_source_line(sec.get("source", "")))
            elif kind == "driver_tree":
                parts.append(_driver_tree_html(sec)); parts.append(_source_line(sec.get("source", "")))
            elif kind == "positioning":
                parts.append(_positioning_html(sec)); parts.append(_source_line(sec.get("source", "")))
            elif kind == "heatmap":
                parts.append(_heatmap_html(sec)); parts.append(_source_line(sec.get("source", "")))
            elif kind == "revenue_build":
                parts.append(_revenue_build_html(sec)); parts.append(_source_line(sec.get("source", "")))
            parts.append(_insight_html(sec.get("so_whats", []), sec.get("watchout")))
            if kind not in EXHIBIT_KINDS:
                parts.append(_source_line(sec.get("source", "")))
            parts.append("</section>")
    rows = [f'<tr><td>{esc(src)}</td><td><span class="cf cf-{conf_}">{conf_}</span></td><td>{esc(note)}</td></tr>'
            for (src, conf_, note) in content["sources"]]
    # Sources close the Appendix chapter when it has sections, else are their own.
    if appendix["sections"]:
        src_head = (f'<h3 style="color:var(--brand);margin:22px 0 6px">{len(chapters)}.'
                    f'{len(appendix["sections"]) + 1} Sources &amp; confidence</h3>')
    else:
        src_head = (f'<h2 style="margin-top:36px;padding-top:18px;border-top:2px solid var(--brand)">'
                    f'<span class="num">{len(chapters) + 1}.</span>Sources &amp; confidence</h2>')
    parts.append('<section>' + src_head +
                 '<table><thead><tr><th>Source</th><th>Conf.</th><th>Note</th></tr></thead><tbody>' + "".join(rows) + '</tbody></table></section>')
    parts.append('<div class="footer"><span class="logo" aria-hidden="true"></span><span>Prepared with Aqmen</span>'
                 f'<span class="spacer">Confidential · {esc(content.get("date", YEAR))}</span></div>')
    parts.append("</div>")
    chart_js = "\n".join(f"aqChart({json.dumps(cid)}, {json.dumps(opt)});" for cid, opt in charts)
    return (head + "\n<body>\n" + "\n".join(parts) + "\n<script>\n" + helpers + "\n\n" + chart_js + "\n\n" + scenario_js + "\n</script>\n</body>\n</html>\n")


# --------------------------------------------------------------------------- #
# Word executive summary
# --------------------------------------------------------------------------- #
def build_summary_docx(content, path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Emu

    FONT, NAVY, BLUE, AZURE, INK, GREY, FILL, WHITE = ("Montserrat", "03045E", "0728A3", "2E6FD6",
                                                       "302E2E", "8A8A8A", "D6EEF5", "FFFFFF")
    WIDTH = 9411

    def run(p, text, size=None, color=None, bold=None, italic=None):
        r = p.add_run(text)
        r.font.name = FONT
        rPr = r._r.get_or_add_rPr()
        rf = rPr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rPr.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), FONT)
        if size: r.font.size = Pt(size / 2)
        if color: r.font.color.rgb = RGBColor.from_string(color)
        if bold is not None: r.font.bold = bold
        if italic is not None: r.font.italic = italic
        return r

    def para(text="", size=None, color=None, bold=None, before=0, after=120, style=None):
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_before = Pt(before / 20)
        p.paragraph_format.space_after = Pt(after / 20)
        if text:
            run(p, text, size, color, bold)
        return p

    def border(p, side, color, sz=8, space=2):
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:space"), str(space)); el.set(qn("w:color"), color)
        pBdr.append(el); pPr.insert(0, pBdr)

    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
        tcPr.append(shd)

    def borders(cell, color=NAVY):
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        for s in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{s}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), color)
            b.append(el)
        tcPr.append(b)

    def margins(cell, v=80, h=120):
        tcPr = cell._tc.get_or_add_tcPr()
        m = OxmlElement("w:tcMar")
        for s, val in (("top", v), ("left", h), ("bottom", v), ("right", h)):
            el = OxmlElement(f"w:{s}"); el.set(qn("w:w"), str(val)); el.set(qn("w:type"), "dxa"); m.append(el)
        tcPr.append(m)

    def table(rows, cols, widths):
        t = doc.add_table(rows=rows, cols=cols)
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))
        for row in t.rows:
            trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement("w:cantSplit"))
            for c, w in zip(row.cells, widths):
                c.width = Emu(w * 635); borders(c)
        return t

    def cell_text(cell, items, size=18, color=None, bold=None):
        first = True
        for it in items:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.paragraph_format.space_after = Pt(2)
            run(p, it, size, color, bold)

    doc = Document(str(ASSETS / "summary-template.docx"))
    # header / footer
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    for r in list(hp.runs):
        r._r.getparent().remove(r._r)
    run(hp, f"AQMEN · {content['title']}", 15, NAVY, True); hp.add_run().add_tab()
    run(hp, "Strictly private & confidential", 15, GREY)
    pPr = hp._p.get_or_add_pPr()
    tabs = pPr.find(qn("w:tabs"))
    if tabs is not None: pPr.remove(tabs)
    tabs = OxmlElement("w:tabs")
    for pos in ("4680", "9360"):
        tb = OxmlElement("w:tab"); tb.set(qn("w:val"), "clear"); tb.set(qn("w:pos"), pos); tabs.append(tb)
    tb = OxmlElement("w:tab"); tb.set(qn("w:val"), "right"); tb.set(qn("w:pos"), str(WIDTH)); tabs.append(tb)
    anchor = pPr.find(qn("w:pBdr"))
    if anchor is None:
        anchor = pPr.find(qn("w:pStyle"))
    if anchor is not None:
        anchor.addnext(tabs)
    else:
        pPr.insert(0, tabs)
    fp = sec.footer.paragraphs[0]
    for r in fp.runs:
        if r.text.strip() and not r._r.findall(qn("w:fldChar")) and not r._r.findall(qn("w:instrText")):
            r.text = "Aqmen AI Limited | Executive summary"; break

    # Title block
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    logo = ASSETS / "logo_deepblue.png"
    if logo.exists():
        p.add_run().add_picture(str(logo), width=Emu(1500000))
    para(content["doctype"].upper() + " · EXECUTIVE SUMMARY", 18, GREY, True, after=80)
    para(content["title"], 40, NAVY, True, after=60)
    if content.get("subtitle"):
        para(content["subtitle"], 22, BLUE, True, after=80)
    meta = f"{content.get('client', '')} · {content.get('date', YEAR)} · Overall confidence {content.get('confidence', 4)}/5".strip(" ·")
    p = para(meta, 16, GREY, after=200); border(p, "bottom", AZURE)

    # Bottom line box
    t = table(1, 1, [WIDTH]); c = t.cell(0, 0); shade(c, FILL); margins(c, 120, 160)
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(4)
    run(p, "BOTTOM LINE  ", 16, BLUE, True); run(p, content["bottom_line"], 20, INK, True)
    para("", after=120)

    # KPIs
    kpis = content.get("kpis", [])
    if kpis:
        n = len(kpis); w = WIDTH // n
        t = table(1, n, [w] * n)
        for j, k in enumerate(kpis):
            c = t.cell(0, j); margins(c, 100, 120)
            val, lbl = k[0], k[1]
            delta = k[2] if len(k) > 2 else None
            cell_text(c, [val], 30, NAVY, True)
            p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0); run(p, str(lbl).upper(), 13, GREY, True)
            if delta:
                p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0); run(p, str(delta), 14, AZURE, True)
        para("", after=120)

    # Executive summary bands
    p = para("Executive summary", 26, NAVY, True, before=200, after=120); border(p, "bottom", AZURE)
    rows = content["exec_summary"]
    t = table(len(rows), 2, [2000, WIDTH - 2000])
    for i, (label, bullets) in enumerate(rows):
        c = t.cell(i, 0); shade(c, NAVY); margins(c, 100, 120); cell_text(c, [label], 18, WHITE, True)
        c = t.cell(i, 1); margins(c, 100, 140)
        cell_text(c, ["•  " + b for b in bullets], 18, INK)
    para("", after=120)

    # Findings by part: the headline of each exhibit, so the reader gets the
    # storyline in the order of the deck. Headlines only — the so-whats live on
    # the slides; here they would push the leave-behind past two pages.
    p = para("Key findings", 26, NAVY, True, before=200, after=120); border(p, "bottom", AZURE)
    for part in split_appendix(content)[0]:
        para(part["name"].upper(), 17, BLUE, True, before=120, after=40)
        shown = 0
        for s in part["sections"]:
            if s["kind"] == "scenario" or s["title"].lower().startswith("appendix"):
                continue
            if shown >= 5:  # the leave-behind is two pages; the deck has the rest
                break
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            run(p, s["headline"], 17, INK)
            shown += 1

    # What you need to believe (optional top-level list)
    if content.get("need_to_believe"):
        p = para("What you need to believe", 26, NAVY, True, before=200, after=120); border(p, "bottom", AZURE)
        for item in content["need_to_believe"]:
            p = doc.add_paragraph(style="List Bullet"); run(p, item, 19, INK)

    # Sources & confidence
    p = para("Sources & confidence", 26, NAVY, True, before=200, after=120); border(p, "bottom", AZURE)
    src = content["sources"]
    t = table(len(src) + 1, 3, [4000, 1200, WIDTH - 5200])
    for j, h in enumerate(("SOURCE", "CONF.", "NOTE")):
        c = t.cell(0, j); shade(c, NAVY); margins(c, 60, 100); cell_text(c, [h], 15, WHITE, True)
        c.paragraphs[0].paragraph_format.keep_with_next = True  # never orphan the header
    p.paragraph_format.keep_with_next = True  # the heading stays with its table
    for i, (s, conf_, note) in enumerate(src, 1):
        for j, v in enumerate((s, f"{conf_}/5", note)):
            c = t.cell(i, j); margins(c, 30, 100); cell_text(c, [str(v)], 15, INK, True if j == 0 else None)
    para("Confidence: 5 primary · 4 credible secondary · 3 triangulated (news cap) · 2 weak · 1 estimate", 14, GREY, before=80)

    doc.core_properties.title = content["title"] + " — Executive summary"
    doc.core_properties.author = "Aqmen AI"
    doc.save(path)
    return path


def docx_to_pdf(docx_path, pdf_path):
    """Windows-only convenience: export through Word COM. Returns True on success."""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        return _docx_to_pdf_powershell(docx_path, pdf_path)
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        d = word.Documents.Open(str(docx_path), False, True)
        d.SaveAs2(str(pdf_path), 17)
        d.Close(False)
    finally:
        word.Quit()
    return os.path.exists(pdf_path)


def _docx_to_pdf_powershell(docx_path, pdf_path):
    import subprocess
    cmd = (f'$w = New-Object -ComObject Word.Application; $w.Visible = $false; '
           f'$d = $w.Documents.Open("{docx_path}", $false, $true); $d.SaveAs2([ref]"{pdf_path}", [ref]17); '
           f'$d.Close($false); $w.Quit()')
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command", cmd], check=True, timeout=180,
                       capture_output=True)
    except Exception:
        return False
    return os.path.exists(pdf_path)
